import os
import sys
from decimal import Decimal
from datetime import datetime
import psycopg2

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.ticker as ticker

from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

DB_CONFIG = {
    'dbname': 'neondb',
    'user': 'neondb_owner',
    'password': 'npg_bFfmWO0xN5vY',
    'host': 'ep-quiet-sky-an30vpqk-pooler.c-6.us-east-1.aws.neon.tech',
    'port': 5432,
    'sslmode': 'require'
}

MESES_NOMES = {
    '2026-06': 'Junho/2026',
    '2026-07': 'Julho/2026',
    '2026-08': 'Agosto/2026',
    '2026-09': 'Setembro/2026',
    '2026-10': 'Outubro/2026*'
}

def format_currency_br(val):
    if val is None:
        return "R$ 0,00"
    v = float(val)
    return f"R$ {v:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.')

def format_number_br(val, decimals=2):
    if val is None:
        return "0"
    v = float(val)
    if decimals == 0:
        return f"{int(round(v)):,}".replace(',', '.')
    fmt = f"{{:,.{decimals}f}}"
    return fmt.format(v).replace(',', 'X').replace('.', ',').replace('X', '.')

def fetch_all_data():
    conn = psycopg2.connect(**DB_CONFIG)
    cur = conn.cursor()

    # 1. Resumo mensal geral
    cur.execute('''
        SELECT 
            TO_CHAR(data, 'YYYY-MM') as mes,
            COUNT(*) as total_abastecimentos,
            ROUND(SUM(quantidade_litros)::numeric, 2) as litros,
            ROUND(SUM(valor_total)::numeric, 2) as valor_total,
            ROUND(AVG(valor_por_litro)::numeric, 3) as preco_medio,
            ROUND((SUM(valor_total) / NULLIF(COUNT(*), 0))::numeric, 2) as ticket_medio,
            ROUND(SUM(CASE WHEN baixa_abastecimento IS TRUE THEN valor_total ELSE 0 END)::numeric, 2) as valor_recebido,
            ROUND(SUM(CASE WHEN baixa_abastecimento IS NOT TRUE THEN valor_total ELSE 0 END)::numeric, 2) as valor_pendente
        FROM abastecimentos
        WHERE deleted_at IS NULL
        GROUP BY TO_CHAR(data, 'YYYY-MM')
        ORDER BY mes ASC;
    ''')
    resumo_mensal = cur.fetchall()

    # 2. Resumo mensal por filial
    cur.execute('''
        SELECT 
            TO_CHAR(data, 'YYYY-MM') as mes,
            COALESCE(local, 'Não Informado') as local,
            COUNT(*) as total_abastecimentos,
            ROUND(SUM(quantidade_litros)::numeric, 2) as litros,
            ROUND(SUM(valor_total)::numeric, 2) as valor_total,
            ROUND(AVG(valor_por_litro)::numeric, 3) as preco_medio,
            ROUND(SUM(CASE WHEN baixa_abastecimento IS TRUE THEN valor_total ELSE 0 END)::numeric, 2) as valor_recebido,
            ROUND(SUM(CASE WHEN baixa_abastecimento IS NOT TRUE THEN valor_total ELSE 0 END)::numeric, 2) as valor_pendente
        FROM abastecimentos
        WHERE deleted_at IS NULL
        GROUP BY TO_CHAR(data, 'YYYY-MM'), local
        ORDER BY mes ASC, local ASC;
    ''')
    resumo_filial = cur.fetchall()

    # 3. Resumo Frota Própria (VIPE) vs Terceiros
    cur.execute('''
        SELECT 
            TO_CHAR(data, 'YYYY-MM') as mes,
            CASE 
                WHEN UPPER(COALESCE(p.nome, a.nome_proprietario, '')) LIKE '%VIPE%' THEN 'Frota Própria (VIPE)'
                ELSE 'Terceiros / Parceiros'
            END as tipo_cliente,
            COUNT(*) as total_abastecimentos,
            ROUND(SUM(a.quantidade_litros)::numeric, 2) as litros,
            ROUND(SUM(a.valor_total)::numeric, 2) as valor_total,
            ROUND(AVG(a.valor_por_litro)::numeric, 3) as preco_medio,
            ROUND(SUM(CASE WHEN a.baixa_abastecimento IS TRUE THEN a.valor_total ELSE 0 END)::numeric, 2) as valor_recebido,
            ROUND(SUM(CASE WHEN a.baixa_abastecimento IS NOT TRUE THEN a.valor_total ELSE 0 END)::numeric, 2) as valor_pendente
        FROM abastecimentos a
        LEFT JOIN proprietarios p ON p.id_proprietario = a.id_proprietario
        WHERE a.deleted_at IS NULL
        GROUP BY 1, 2
        ORDER BY mes ASC, tipo_cliente ASC;
    ''')
    resumo_segmento = cur.fetchall()

    # 4. Top Proprietários Geral
    cur.execute('''
        SELECT 
            COALESCE(p.nome, a.nome_proprietario, 'Não Identificado') as proprietario,
            COUNT(*) as total_abastecimentos,
            ROUND(SUM(a.quantidade_litros)::numeric, 2) as litros,
            ROUND(SUM(a.valor_total)::numeric, 2) as valor_total,
            ROUND(SUM(CASE WHEN a.baixa_abastecimento IS TRUE THEN a.valor_total ELSE 0 END)::numeric, 2) as valor_recebido,
            ROUND(SUM(CASE WHEN a.baixa_abastecimento IS NOT TRUE THEN a.valor_total ELSE 0 END)::numeric, 2) as valor_pendente,
            ROUND(AVG(a.valor_por_litro)::numeric, 3) as preco_medio
        FROM abastecimentos a
        LEFT JOIN proprietarios p ON p.id_proprietario = a.id_proprietario
        WHERE a.deleted_at IS NULL
        GROUP BY COALESCE(p.nome, a.nome_proprietario, 'Não Identificado')
        ORDER BY valor_total DESC;
    ''')
    top_proprietarios = cur.fetchall()

    # 5. Matriz Proprietários x Mês
    cur.execute('''
        SELECT 
            COALESCE(p.nome, a.nome_proprietario, 'Não Identificado') as proprietario,
            TO_CHAR(a.data, 'YYYY-MM') as mes,
            ROUND(SUM(a.valor_total)::numeric, 2) as valor_total,
            ROUND(SUM(a.quantidade_litros)::numeric, 2) as litros
        FROM abastecimentos a
        LEFT JOIN proprietarios p ON p.id_proprietario = a.id_proprietario
        WHERE a.deleted_at IS NULL
        GROUP BY 1, 2
        ORDER BY 1, 2;
    ''')
    proprietarios_mes = cur.fetchall()

    # 6. Base analítica completa
    cur.execute('''
        SELECT 
            a.id_abastecimento,
            COALESCE(a.data_hora, a.data::timestamp) as data_hora,
            TO_CHAR(a.data, 'YYYY-MM') as mes,
            a.local,
            COALESCE(p.nome, a.nome_proprietario, '—') as proprietario,
            COALESCE(v.placa, a.placa1, '—') as placa,
            COALESCE(v.modelo, '—') as modelo_veiculo,
            COALESCE(m.nome, a.nome_motorista, '—') as motorista,
            a.tipo_combustivel,
            a.quantidade_litros::numeric as litros,
            a.valor_por_litro::numeric as valor_litro,
            a.valor_total::numeric as valor_total,
            CASE WHEN a.baixa_abastecimento IS TRUE THEN 'Baixado' ELSE 'Pendente' END as status_baixa
        FROM abastecimentos a
        LEFT JOIN proprietarios p ON p.id_proprietario = a.id_proprietario
        LEFT JOIN veiculos v ON v.id_veiculo = a.id_veiculo
        LEFT JOIN motoristas m ON m.id_motorista = a.id_motorista
        WHERE a.deleted_at IS NULL
        ORDER BY COALESCE(a.data_hora, a.data::timestamp) DESC;
    ''')
    base_detalhada = cur.fetchall()

    conn.close()
    return {
        'resumo_mensal': resumo_mensal,
        'resumo_filial': resumo_filial,
        'resumo_segmento': resumo_segmento,
        'top_proprietarios': top_proprietarios,
        'proprietarios_mes': proprietarios_mes,
        'base_detalhada': base_detalhada
    }

# ==============================================================================
# GERAÇÃO DO EXCEL (XLSX)
# ==============================================================================
def gerar_excel(data, output_path):
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # Styles
    navy_dark = "1B365D"
    navy_header = "2C3E50"
    blue_accent = "0284C7"
    gray_light = "F8FAFC"
    gray_border = "CBD5E1"
    green_soft = "ECFDF5"
    amber_soft = "FFFBEB"

    font_title = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="E2E8F0")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_sub_header = Font(name="Calibri", size=10, bold=True, color="1E293B")
    font_bold = Font(name="Calibri", size=11, bold=True, color="0F172A")
    font_regular = Font(name="Calibri", size=10, color="1E293B")
    font_kpi_label = Font(name="Calibri", size=9, bold=True, color="64748B")
    font_kpi_val = Font(name="Calibri", size=14, bold=True, color="1B365D")

    fill_title = PatternFill(start_color=navy_dark, end_color=navy_dark, fill_type="solid")
    fill_header = PatternFill(start_color=navy_header, end_color=navy_header, fill_type="solid")
    fill_zebra = PatternFill(start_color=gray_light, end_color=gray_light, fill_type="solid")
    fill_total = PatternFill(start_color="E2E8F0", end_color="E2E8F0", fill_type="solid")
    fill_kpi = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    border_thin = Border(
        left=Side(style='thin', color=gray_border),
        right=Side(style='thin', color=gray_border),
        top=Side(style='thin', color=gray_border),
        bottom=Side(style='thin', color=gray_border)
    )
    border_total = Border(
        top=Side(style='thin', color="0F172A"),
        bottom=Side(style='double', color="0F172A")
    )

    align_center = Alignment(horizontal='center', vertical='center')
    align_left = Alignment(horizontal='left', vertical='center')
    align_right = Alignment(horizontal='right', vertical='center')

    num_fmt_curr = 'R$ #,##0.00'
    num_fmt_lit = '#,##0.00'
    num_fmt_int = '#,##0'
    num_fmt_pct = '0.0%'
    num_fmt_price = 'R$ #,##0.000'

    # --------------------------------------------------------------------------
    # ABA 1: RESUMO EXECUTIVO MENSAL
    # --------------------------------------------------------------------------
    ws1 = wb.create_sheet(title="Resumo Mensal")
    ws1.views.sheetView[0].showGridLines = True

    # Title Banner
    ws1.merge_cells("A1:J1")
    ws1["A1"] = "RELATÓRIO DE FATURAMENTO MENSAL — SISTEMA DE ABASTECIMENTO"
    ws1["A1"].font = font_title
    ws1["A1"].fill = fill_title
    ws1["A1"].alignment = align_center
    ws1.row_dimensions[1].height = 36

    ws1.merge_cells("A2:J2")
    ws1["A2"] = f"Extraído em: {datetime.now().strftime('%d/%m/%Y %H:%M')} | Período: Junho/2026 a Outubro/2026 | Fonte: Base de Dados Operacional"
    ws1["A2"].font = font_subtitle
    ws1["A2"].fill = PatternFill(start_color="334155", end_color="334155", fill_type="solid")
    ws1["A2"].alignment = align_center
    ws1.row_dimensions[2].height = 20

    # Total KPIs Cards
    tot_abast = sum(r[1] for r in data['resumo_mensal'])
    tot_lit = sum(float(r[2]) for r in data['resumo_mensal'])
    tot_val = sum(float(r[3]) for r in data['resumo_mensal'])
    tot_rec = sum(float(r[6]) for r in data['resumo_mensal'])
    tot_pend = sum(float(r[7]) for r in data['resumo_mensal'])
    preco_medio_global = tot_val / tot_lit if tot_lit else 0

    kpis = [
        ("A4:B4", "A5:B5", "FATURAMENTO TOTAL", tot_val, num_fmt_curr),
        ("C4:D4", "C5:D5", "VOLUME TOTAL (L)", tot_lit, num_fmt_lit),
        ("E4:F4", "E5:F5", "TOTAL ABASTECIMENTOS", tot_abast, num_fmt_int),
        ("G4:H4", "G5:H5", "PREÇO MÉDIO / L", preco_medio_global, num_fmt_price),
        ("I4:J4", "I5:J5", "RECEBIDO / BAIXADO", tot_rec, num_fmt_curr),
    ]
    for top_range, val_range, lbl, val, fmt in kpis:
        ws1.merge_cells(top_range)
        ws1.merge_cells(val_range)
        top_cell = ws1[top_range.split(":")[0]]
        val_cell = ws1[val_range.split(":")[0]]
        top_cell.value = lbl
        top_cell.font = font_kpi_label
        top_cell.alignment = align_center
        top_cell.fill = fill_kpi
        val_cell.value = val
        val_cell.font = font_kpi_val
        val_cell.alignment = align_center
        val_cell.fill = fill_kpi
        val_cell.number_format = fmt

    for r in range(4, 6):
        ws1.row_dimensions[r].height = 22

    # Table Header
    headers1 = [
        "Mês", "Período", "Abastecimentos", "Volume (Litros)", "Faturamento (R$)",
        "Preço Médio/L", "Ticket Médio", "Recebido (R$)", "Pendente (R$)", "% Recebido"
    ]
    start_row = 7
    ws1.row_dimensions[start_row].height = 26
    for col_idx, h in enumerate(headers1, 1):
        cell = ws1.cell(row=start_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_thin

    curr_row = start_row + 1
    for idx, r in enumerate(data['resumo_mensal']):
        mes_cod = r[0]
        mes_nome = MESES_NOMES.get(mes_cod, mes_cod)
        qtd = r[1]
        litros = float(r[2])
        valor = float(r[3])
        p_medio = float(r[4])
        t_medio = float(r[5])
        rec = float(r[6])
        pend = float(r[7])
        pct_rec = rec / valor if valor > 0 else 0

        row_fill = fill_zebra if idx % 2 == 1 else None

        values = [
            (mes_cod, align_center, '@'),
            (mes_nome, align_left, '@'),
            (qtd, align_right, num_fmt_int),
            (litros, align_right, num_fmt_lit),
            (valor, align_right, num_fmt_curr),
            (p_medio, align_right, num_fmt_price),
            (t_medio, align_right, num_fmt_curr),
            (rec, align_right, num_fmt_curr),
            (pend, align_right, num_fmt_curr),
            (pct_rec, align_right, num_fmt_pct),
        ]

        ws1.row_dimensions[curr_row].height = 20
        for col_idx, (v, aln, fmt) in enumerate(values, 1):
            cell = ws1.cell(row=curr_row, column=col_idx, value=v)
            cell.font = font_regular
            cell.alignment = aln
            cell.border = border_thin
            if row_fill:
                cell.fill = row_fill
            cell.number_format = fmt

        curr_row += 1

    # Totals Row
    ws1.row_dimensions[curr_row].height = 24
    totals1 = [
        ("TOTAL", align_center, '@'),
        ("Consolidado", align_left, '@'),
        (tot_abast, align_right, num_fmt_int),
        (tot_lit, align_right, num_fmt_lit),
        (tot_val, align_right, num_fmt_curr),
        (preco_medio_global, align_right, num_fmt_price),
        (tot_val / tot_abast if tot_abast else 0, align_right, num_fmt_curr),
        (tot_rec, align_right, num_fmt_curr),
        (tot_pend, align_right, num_fmt_curr),
        (tot_rec / tot_val if tot_val > 0 else 0, align_right, num_fmt_pct),
    ]
    for col_idx, (v, aln, fmt) in enumerate(totals1, 1):
        cell = ws1.cell(row=curr_row, column=col_idx, value=v)
        cell.font = font_bold
        cell.alignment = aln
        cell.fill = fill_total
        cell.border = border_total
        cell.number_format = fmt

    # --------------------------------------------------------------------------
    # ABA 2: FATURAMENTO POR FILIAL E MÊS
    # --------------------------------------------------------------------------
    ws2 = wb.create_sheet(title="Por Filial")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:H1")
    ws2["A1"] = "FATURAMENTO MENSAL POR FILIAL (MATRIZ vs VIANA)"
    ws2["A1"].font = font_title
    ws2["A1"].fill = fill_title
    ws2["A1"].alignment = align_center
    ws2.row_dimensions[1].height = 36

    headers2 = ["Mês", "Filial / Local", "Abastecimentos", "Volume (Litros)", "Faturamento (R$)", "Preço Médio/L", "Recebido (R$)", "Pendente (R$)"]
    ws2.row_dimensions[3].height = 24
    for col_idx, h in enumerate(headers2, 1):
        cell = ws2.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_thin

    curr_row = 4
    for idx, r in enumerate(data['resumo_filial']):
        mes_cod = r[0]
        local = r[1]
        qtd = r[2]
        litros = float(r[3])
        valor = float(r[4])
        p_medio = float(r[5])
        rec = float(r[6])
        pend = float(r[7])

        row_fill = fill_zebra if idx % 2 == 1 else None

        values = [
            (MESES_NOMES.get(mes_cod, mes_cod), align_left, '@'),
            (local, align_center, '@'),
            (qtd, align_right, num_fmt_int),
            (litros, align_right, num_fmt_lit),
            (valor, align_right, num_fmt_curr),
            (p_medio, align_right, num_fmt_price),
            (rec, align_right, num_fmt_curr),
            (pend, align_right, num_fmt_curr),
        ]
        ws2.row_dimensions[curr_row].height = 20
        for col_idx, (v, aln, fmt) in enumerate(values, 1):
            cell = ws2.cell(row=curr_row, column=col_idx, value=v)
            cell.font = font_regular
            cell.alignment = aln
            cell.border = border_thin
            if row_fill:
                cell.fill = row_fill
            cell.number_format = fmt
        curr_row += 1

    # Totals Row
    ws2.row_dimensions[curr_row].height = 24
    totals2 = [
        ("TOTAL GERAL", align_left, '@'),
        ("Todas", align_center, '@'),
        (tot_abast, align_right, num_fmt_int),
        (tot_lit, align_right, num_fmt_lit),
        (tot_val, align_right, num_fmt_curr),
        (preco_medio_global, align_right, num_fmt_price),
        (tot_rec, align_right, num_fmt_curr),
        (tot_pend, align_right, num_fmt_curr),
    ]
    for col_idx, (v, aln, fmt) in enumerate(totals2, 1):
        cell = ws2.cell(row=curr_row, column=col_idx, value=v)
        cell.font = font_bold
        cell.alignment = aln
        cell.fill = fill_total
        cell.border = border_total
        cell.number_format = fmt

    # --------------------------------------------------------------------------
    # ABA 3: FROTA PRÓPRIA VS TERCEIROS
    # --------------------------------------------------------------------------
    ws3 = wb.create_sheet(title="Frota vs Terceiros")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:H1")
    ws3["A1"] = "FATURAMENTO POR SEGMENTO — FROTA PRÓPRIA (VIPE) vs TERCEIROS"
    ws3["A1"].font = font_title
    ws3["A1"].fill = fill_title
    ws3["A1"].alignment = align_center
    ws3.row_dimensions[1].height = 36

    headers3 = ["Mês", "Segmento", "Abastecimentos", "Volume (Litros)", "Faturamento (R$)", "Preço Médio/L", "Recebido (R$)", "Pendente (R$)"]
    ws3.row_dimensions[3].height = 24
    for col_idx, h in enumerate(headers3, 1):
        cell = ws3.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_thin

    curr_row = 4
    for idx, r in enumerate(data['resumo_segmento']):
        mes_cod = r[0]
        segmento = r[1]
        qtd = r[2]
        litros = float(r[3])
        valor = float(r[4])
        p_medio = float(r[5])
        rec = float(r[6])
        pend = float(r[7])

        row_fill = fill_zebra if idx % 2 == 1 else None

        values = [
            (MESES_NOMES.get(mes_cod, mes_cod), align_left, '@'),
            (segmento, align_left, '@'),
            (qtd, align_right, num_fmt_int),
            (litros, align_right, num_fmt_lit),
            (valor, align_right, num_fmt_curr),
            (p_medio, align_right, num_fmt_price),
            (rec, align_right, num_fmt_curr),
            (pend, align_right, num_fmt_curr),
        ]
        ws3.row_dimensions[curr_row].height = 20
        for col_idx, (v, aln, fmt) in enumerate(values, 1):
            cell = ws3.cell(row=curr_row, column=col_idx, value=v)
            cell.font = font_regular
            cell.alignment = aln
            cell.border = border_thin
            if row_fill:
                cell.fill = row_fill
            cell.number_format = fmt
        curr_row += 1

    # Totals Row
    ws3.row_dimensions[curr_row].height = 24
    totals3 = [
        ("TOTAL GERAL", align_left, '@'),
        ("Consolidado", align_left, '@'),
        (tot_abast, align_right, num_fmt_int),
        (tot_lit, align_right, num_fmt_lit),
        (tot_val, align_right, num_fmt_curr),
        (preco_medio_global, align_right, num_fmt_price),
        (tot_rec, align_right, num_fmt_curr),
        (tot_pend, align_right, num_fmt_curr),
    ]
    for col_idx, (v, aln, fmt) in enumerate(totals3, 1):
        cell = ws3.cell(row=curr_row, column=col_idx, value=v)
        cell.font = font_bold
        cell.alignment = aln
        cell.fill = fill_total
        cell.border = border_total
        cell.number_format = fmt

    # --------------------------------------------------------------------------
    # ABA 4: RANKING DE PROPRIETÁRIOS
    # --------------------------------------------------------------------------
    ws4 = wb.create_sheet(title="Por Proprietário")
    ws4.views.sheetView[0].showGridLines = True

    # Build matrix: proprietario -> {mes: valor}
    meses_lista = [r[0] for r in data['resumo_mensal']]
    matriz_prop = {}
    for r in data['proprietarios_mes']:
        p_nome, mes, val, lit = r[0], r[1], float(r[2]), float(r[3])
        if p_nome not in matriz_prop:
            matriz_prop[p_nome] = {}
        matriz_prop[p_nome][mes] = val

    ws4.merge_cells("A1:K1")
    ws4["A1"] = "FATURAMENTO MENSAL POR PROPRIETÁRIO / CLIENTE"
    ws4["A1"].font = font_title
    ws4["A1"].fill = fill_title
    ws4["A1"].alignment = align_center
    ws4.row_dimensions[1].height = 36

    headers4 = ["Proprietário", "Abastecimentos", "Volume (L)"] + [MESES_NOMES.get(m, m) for m in meses_lista] + ["Total Faturado (R$)", "Recebido (R$)", "Pendente (R$)"]
    ws4.row_dimensions[3].height = 24
    for col_idx, h in enumerate(headers4, 1):
        cell = ws4.cell(row=3, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_thin

    curr_row = 4
    for idx, r in enumerate(data['top_proprietarios']):
        p_nome = r[0]
        qtd = r[1]
        litros = float(r[2])
        val_tot = float(r[3])
        rec = float(r[4])
        pend = float(r[5])

        row_fill = fill_zebra if idx % 2 == 1 else None

        row_vals = [
            (p_nome, align_left, '@'),
            (qtd, align_right, num_fmt_int),
            (litros, align_right, num_fmt_lit),
        ]
        # Meses
        p_meses = matriz_prop.get(p_nome, {})
        for m in meses_lista:
            v_m = p_meses.get(m, 0.0)
            row_vals.append((v_m, align_right, num_fmt_curr))

        row_vals.extend([
            (val_tot, align_right, num_fmt_curr),
            (rec, align_right, num_fmt_curr),
            (pend, align_right, num_fmt_curr),
        ])

        ws4.row_dimensions[curr_row].height = 20
        for col_idx, (v, aln, fmt) in enumerate(row_vals, 1):
            cell = ws4.cell(row=curr_row, column=col_idx, value=v)
            cell.font = font_regular
            cell.alignment = aln
            cell.border = border_thin
            if row_fill:
                cell.fill = row_fill
            cell.number_format = fmt
        curr_row += 1

    # --------------------------------------------------------------------------
    # ABA 5: BASE DETALHADA ANALÍTICA (4.396 registros)
    # --------------------------------------------------------------------------
    ws5 = wb.create_sheet(title="Base Detalhada")
    ws5.views.sheetView[0].showGridLines = True

    headers5 = [
        "ID Abastecimento", "Data / Hora", "Mês", "Filial", "Proprietário",
        "Placa", "Modelo Veículo", "Motorista", "Combustível",
        "Volume (L)", "Valor / Litro (R$)", "Valor Total (R$)", "Status Baixa"
    ]
    ws5.row_dimensions[1].height = 26
    for col_idx, h in enumerate(headers5, 1):
        cell = ws5.cell(row=1, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_thin

    curr_row = 2
    for idx, r in enumerate(data['base_detalhada']):
        id_abast = str(r[0])
        dt_hora = r[1].strftime('%d/%m/%Y %H:%M') if r[1] else ''
        mes = r[2]
        filial = r[3] or ''
        prop = r[4] or ''
        placa = r[5] or ''
        modelo = r[6] or ''
        motorista = r[7] or ''
        comb = r[8] or ''
        lit = float(r[9]) if r[9] is not None else 0.0
        v_lit = float(r[10]) if r[10] is not None else 0.0
        v_tot = float(r[11]) if r[11] is not None else 0.0
        status_bx = r[12] or ''

        row_fill = fill_zebra if idx % 2 == 1 else None

        row_vals = [
            (id_abast, align_center, '@'),
            (dt_hora, align_center, '@'),
            (mes, align_center, '@'),
            (filial, align_center, '@'),
            (prop, align_left, '@'),
            (placa, align_center, '@'),
            (modelo, align_left, '@'),
            (motorista, align_left, '@'),
            (comb, align_center, '@'),
            (lit, align_right, num_fmt_lit),
            (v_lit, align_right, num_fmt_price),
            (v_tot, align_right, num_fmt_curr),
            (status_bx, align_center, '@'),
        ]

        ws5.row_dimensions[curr_row].height = 19
        for col_idx, (v, aln, fmt) in enumerate(row_vals, 1):
            cell = ws5.cell(row=curr_row, column=col_idx, value=v)
            cell.font = font_regular
            cell.alignment = aln
            cell.border = border_thin
            if row_fill:
                cell.fill = row_fill
            cell.number_format = fmt
        curr_row += 1

    # Auto-adjust column widths for all sheets
    for ws in [ws1, ws2, ws3, ws4, ws5]:
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                # ignore merged cells in length calculation
                if cell.row in [1, 2] and ws != ws5:
                    continue
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Freeze panes
    ws1.freeze_panes = "A8"
    ws2.freeze_panes = "A4"
    ws3.freeze_panes = "A4"
    ws4.freeze_panes = "A4"
    ws5.freeze_panes = "A2"

    # Enable autofilter on base detalhada
    ws5.auto_filter.ref = f"A1:M{curr_row - 1}"

    wb.save(output_path)
    print(f"Excel salvo com sucesso em: {output_path}")


# ==============================================================================
# GERAÇÃO DOS GRÁFICOS (MATPLOTLIB)
# ==============================================================================
def gerar_graficos(data, charts_dir):
    os.makedirs(charts_dir, exist_ok=True)
    plt.rcParams['font.sans-serif'] = 'DejaVu Sans'
    plt.rcParams['axes.edgecolor'] = '#CBD5E1'
    plt.rcParams['axes.linewidth'] = 0.8

    # 1. Gráfico de Faturamento Mensal por Filial
    fig, ax = plt.subplots(figsize=(8.5, 3.4), dpi=200)
    meses = [r[0] for r in data['resumo_mensal']]
    meses_labels = [MESES_NOMES.get(m, m).replace('/2026', '') for m in meses]

    # Map valores por filial
    val_matriz = []
    val_viana = []
    filial_dict = {(r[0], r[1]): float(r[4]) for r in data['resumo_filial']}
    for m in meses:
        val_matriz.append(filial_dict.get((m, 'Matriz'), 0.0) / 1000.0) # em mil R$
        val_viana.append(filial_dict.get((m, 'Viana'), 0.0) / 1000.0)

    import numpy as np
    x = np.arange(len(meses))
    width = 0.36

    bars1 = ax.bar(x - width/2, val_matriz, width, label='Matriz', color='#1B365D', edgecolor='#0F172A')
    bars2 = ax.bar(x + width/2, val_viana, width, label='Viana', color='#0284C7', edgecolor='#0369A1')

    ax.set_title('Faturamento Mensal por Filial (em Mil R$)', fontsize=12, fontweight='bold', color='#1E293B', pad=12)
    ax.set_xticks(x)
    ax.set_xticklabels(meses_labels, fontsize=10, fontweight='bold', color='#334155')
    ax.set_ylabel('Mil R$', fontsize=10, color='#64748B')
    ax.yaxis.set_major_formatter(ticker.FuncFormatter(lambda v, p: f"R$ {v:,.0f}k"))
    ax.grid(axis='y', linestyle='--', alpha=0.5, color='#CBD5E1')
    ax.set_axisbelow(True)
    ax.legend(frameon=True, facecolor='#F8FAFC', edgecolor='#CBD5E1')

    for bar in bars1:
        h = bar.get_height()
        if h > 0:
            ax.annotate(f"R$ {h:,.0f}k",
                        xy=(bar.get_x() + bar.get_width() / 2, h),
                        xytext=(0, 3), textcoords="offset points",
                        ha='center', va='bottom', fontsize=7.5, fontweight='bold', color='#1B365D')

    for bar in bars2:
        h = bar.get_height()
        if h > 0:
            ax.annotate(f"R$ {h:,.0f}k",
                        xy=(bar.get_x() + bar.get_width() / 2, h),
                        xytext=(0, 3), textcoords="offset points",
                        ha='center', va='bottom', fontsize=7.5, fontweight='bold', color='#0284C7')

    plt.tight_layout()
    chart1_path = os.path.join(charts_dir, "faturamento_filiais.png")
    plt.savefig(chart1_path, dpi=200)
    plt.close()

    # 2. Gráfico de Donut: Segmento e Status de Cobrança
    tot_val = sum(float(r[3]) for r in data['resumo_mensal'])
    fig, (ax1, ax2) = plt.subplots(1, 2, figsize=(8.5, 3.2), dpi=200)

    # Subplot 1: Frota Própria vs Terceiros
    tot_propria = sum(float(r[4]) for r in data['resumo_segmento'] if 'VIPE' in r[1])
    tot_terceiros = sum(float(r[4]) for r in data['resumo_segmento'] if 'VIPE' not in r[1])
    wedges1, texts1, autotexts1 = ax1.pie(
        [tot_propria, tot_terceiros],
        autopct='%1.1f%%',
        pctdistance=0.76,
        startangle=140,
        colors=['#1B365D', '#0284C7'],
        wedgeprops=dict(width=0.48, edgecolor='w', linewidth=2),
    )
    for at in autotexts1:
        at.set_color('white')
        at.set_fontsize(8.5)
        at.set_fontweight('bold')
    ax1.set_title('Distribuição por Segmento', fontsize=11, fontweight='bold', color='#1E293B', pad=8)
    ax1.legend(
        wedges1, 
        [f"Frota Própria ({tot_propria/tot_val*100:.1f}%)", f"Terceiros ({tot_terceiros/tot_val*100:.1f}%)"],
        loc='lower center',
        bbox_to_anchor=(0.5, -0.16),
        frameon=False,
        fontsize=8.5,
        ncol=2
    )

    # Subplot 2: Recebido vs Pendente
    tot_rec = sum(float(r[6]) for r in data['resumo_mensal'])
    tot_pend = sum(float(r[7]) for r in data['resumo_mensal'])
    wedges2, texts2, autotexts2 = ax2.pie(
        [tot_rec, tot_pend],
        autopct='%1.1f%%',
        pctdistance=0.76,
        startangle=140,
        colors=['#10B981', '#F59E0B'],
        wedgeprops=dict(width=0.48, edgecolor='w', linewidth=2),
    )
    for at in autotexts2:
        at.set_color('white')
        at.set_fontsize(8.5)
        at.set_fontweight('bold')
    ax2.set_title('Status Financeiro (Baixas)', fontsize=11, fontweight='bold', color='#1E293B', pad=8)
    ax2.legend(
        wedges2,
        [f"Baixado ({tot_rec/tot_val*100:.1f}%)", f"Pendente ({tot_pend/tot_val*100:.1f}%)"],
        loc='lower center',
        bbox_to_anchor=(0.5, -0.16),
        frameon=False,
        fontsize=8.5,
        ncol=2
    )

    plt.tight_layout()
    chart2_path = os.path.join(charts_dir, "segmentos_baixas.png")
    plt.savefig(chart2_path, dpi=200, bbox_inches='tight')
    plt.close()

    return chart1_path, chart2_path


# ==============================================================================
# GERAÇÃO DO PDF (REPORTLAB - LANDSCAPE EXECUTIVE)
# ==============================================================================
class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            canvas.Canvas.showPage(self)
        canvas.Canvas.save(self)

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        
        # Header line (pages > 1)
        if self._pageNumber > 1:
            self.drawString(36, 565, "VIPE TRANSPORTES — Relatório Executivo de Faturamento Mensal")
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(36, 558, 806, 558)

        # Footer line
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(36, 32, 806, 32)

        data_geracao = datetime.now().strftime("%d/%m/%Y às %H:%M")
        self.drawString(36, 20, f"Gerado em {data_geracao} | Sistema de Gestão de Abastecimento | Documento Confidencial")
        page_text = f"Página {self._pageNumber} de {page_count}"
        self.drawRightString(806, 20, page_text)
        self.restoreState()


def gerar_pdf(data, output_path, charts_dir):
    chart1_path, chart2_path = gerar_graficos(data, charts_dir)

    doc = SimpleDocTemplate(
        output_path,
        pagesize=landscape(A4),
        leftMargin=36,
        rightMargin=36,
        topMargin=40,
        bottomMargin=42
    )

    styles = getSampleStyleSheet()
    
    # Custom Styles
    style_title = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1B365D')
    )
    style_subtitle = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor('#64748B')
    )
    style_section = ParagraphStyle(
        'DocSection',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1B365D'),
        spaceAfter=6,
        spaceBefore=10
    )
    style_kpi_lbl = ParagraphStyle(
        'KpiLbl',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#64748B'),
        alignment=1 # Center
    )
    style_kpi_val = ParagraphStyle(
        'KpiVal',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#1B365D'),
        alignment=1 # Center
    )
    style_th = ParagraphStyle(
        'TH',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=1 # Center
    )
    style_td = ParagraphStyle(
        'TD',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1E293B')
    )
    style_td_bold = ParagraphStyle(
        'TDBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#0F172A')
    )
    style_td_right = ParagraphStyle(
        'TDRight',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#1E293B'),
        alignment=2 # Right
    )
    style_td_right_bold = ParagraphStyle(
        'TDRightBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#0F172A'),
        alignment=2 # Right
    )

    story = []

    # Header Banner
    header_data = [
        [
            Paragraph("RELATÓRIO EXECUTIVO DE FATURAMENTO MENSAL", style_title),
            Paragraph(f"<b>Período:</b> Junho a Outubro/2026<br/><b>Emissão:</b> {datetime.now().strftime('%d/%m/%Y %H:%M')}<br/><b>Filiais:</b> Matriz e Viana", style_subtitle)
        ]
    ]
    t_header = Table(header_data, colWidths=[520, 250])
    t_header.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('ALIGN', (1,0), (1,0), 'RIGHT'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ('TOPPADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(t_header)
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1B365D"), spaceAfter=10))

    # KPI Summary Cards
    tot_abast = sum(r[1] for r in data['resumo_mensal'])
    tot_lit = sum(float(r[2]) for r in data['resumo_mensal'])
    tot_val = sum(float(r[3]) for r in data['resumo_mensal'])
    tot_rec = sum(float(r[6]) for r in data['resumo_mensal'])
    tot_pend = sum(float(r[7]) for r in data['resumo_mensal'])
    preco_medio_global = tot_val / tot_lit if tot_lit else 0

    kpi_table_data = [
        [
            Paragraph("FATURAMENTO TOTAL", style_kpi_lbl),
            Paragraph("VOLUME TOTAL", style_kpi_lbl),
            Paragraph("ABASTECIMENTOS", style_kpi_lbl),
            Paragraph("PREÇO MÉDIO / L", style_kpi_lbl),
            Paragraph("TOTAL RECEBIDO", style_kpi_lbl),
            Paragraph("TOTAL PENDENTE", style_kpi_lbl),
        ],
        [
            Paragraph(format_currency_br(tot_val), style_kpi_val),
            Paragraph(f"{format_number_br(tot_lit, 2)} L", style_kpi_val),
            Paragraph(format_number_br(tot_abast, 0), style_kpi_val),
            Paragraph(f"R$ {preco_medio_global:.3f}".replace('.', ','), style_kpi_val),
            Paragraph(format_currency_br(tot_rec), ParagraphStyle('KpiGreen', parent=style_kpi_val, textColor=colors.HexColor('#059669'))),
            Paragraph(format_currency_br(tot_pend), ParagraphStyle('KpiAmber', parent=style_kpi_val, textColor=colors.HexColor('#D97706'))),
        ]
    ]
    t_kpis = Table(kpi_table_data, colWidths=[128, 128, 128, 128, 128, 130])
    t_kpis.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0,0), (-1,0), 6),
        ('BOTTOMPADDING', (0,0), (-1,0), 2),
        ('TOPPADDING', (0,1), (-1,1), 2),
        ('BOTTOMPADDING', (0,1), (-1,1), 6),
    ]))
    story.append(t_kpis)
    story.append(Spacer(1, 10))

    # Tabela 1: Resumo Mensal Consolidado
    story.append(Paragraph("1. Faturamento Mensal Consolidado", style_section))
    t1_data = [
        [
            Paragraph("Mês", style_th),
            Paragraph("Abastecimentos", style_th),
            Paragraph("Volume (Litros)", style_th),
            Paragraph("Faturamento (R$)", style_th),
            Paragraph("Preço Médio/L", style_th),
            Paragraph("Ticket Médio", style_th),
            Paragraph("Recebido / Baixado", style_th),
            Paragraph("Pendente", style_th),
            Paragraph("% Recebido", style_th),
        ]
    ]
    for idx, r in enumerate(data['resumo_mensal']):
        mes_cod = r[0]
        mes_nome = MESES_NOMES.get(mes_cod, mes_cod)
        qtd = r[1]
        litros = float(r[2])
        valor = float(r[3])
        p_medio = float(r[4])
        t_medio = float(r[5])
        rec = float(r[6])
        pend = float(r[7])
        pct_rec = (rec / valor * 100.0) if valor > 0 else 0

        t1_data.append([
            Paragraph(mes_nome, style_td_bold),
            Paragraph(format_number_br(qtd, 0), style_td_right),
            Paragraph(format_number_br(litros, 2), style_td_right),
            Paragraph(format_currency_br(valor), style_td_right_bold),
            Paragraph(f"R$ {p_medio:.3f}".replace('.', ','), style_td_right),
            Paragraph(format_currency_br(t_medio), style_td_right),
            Paragraph(format_currency_br(rec), style_td_right),
            Paragraph(format_currency_br(pend), style_td_right),
            Paragraph(f"{pct_rec:.1f}%".replace('.', ','), style_td_right_bold),
        ])

    # Linha Total
    t1_data.append([
        Paragraph("TOTAL CONSOLIDADO", style_td_bold),
        Paragraph(format_number_br(tot_abast, 0), style_td_right_bold),
        Paragraph(format_number_br(tot_lit, 2), style_td_right_bold),
        Paragraph(format_currency_br(tot_val), style_td_right_bold),
        Paragraph(f"R$ {preco_medio_global:.3f}".replace('.', ','), style_td_right_bold),
        Paragraph(format_currency_br(tot_val / tot_abast), style_td_right_bold),
        Paragraph(format_currency_br(tot_rec), style_td_right_bold),
        Paragraph(format_currency_br(tot_pend), style_td_right_bold),
        Paragraph(f"{(tot_rec/tot_val*100):.1f}%".replace('.', ','), style_td_right_bold),
    ])

    t1 = Table(t1_data, colWidths=[95, 75, 85, 95, 75, 75, 95, 95, 80])
    t1_style = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1B365D')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#E2E8F0')),
    ]
    for r_idx in range(1, len(t1_data) - 1):
        if r_idx % 2 == 0:
            t1_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), colors.HexColor('#F8FAFC')))
    t1.setStyle(TableStyle(t1_style))
    story.append(t1)
    story.append(Spacer(1, 10))

    # Gráficos lado a lado
    chart_table_data = [
        [
            Image(chart1_path, width=380, height=152),
            Image(chart2_path, width=380, height=152)
        ]
    ]
    t_charts = Table(chart_table_data, colWidths=[385, 385])
    t_charts.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ('TOPPADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(t_charts)

    # PAGE BREAK: Página 2 para detalhamentos
    story.append(PageBreak())

    # Tabela 2: Faturamento por Filial
    story.append(Paragraph("2. Faturamento Mensal por Filial (Matriz vs Viana)", style_section))
    t2_data = [
        [
            Paragraph("Mês", style_th),
            Paragraph("Filial", style_th),
            Paragraph("Abastecimentos", style_th),
            Paragraph("Volume (Litros)", style_th),
            Paragraph("Faturamento (R$)", style_th),
            Paragraph("Preço Médio/L", style_th),
            Paragraph("Recebido (R$)", style_th),
            Paragraph("Pendente (R$)", style_th),
        ]
    ]
    for idx, r in enumerate(data['resumo_filial']):
        mes_cod = r[0]
        filial = r[1]
        qtd = r[2]
        litros = float(r[3])
        valor = float(r[4])
        p_medio = float(r[5])
        rec = float(r[6])
        pend = float(r[7])

        t2_data.append([
            Paragraph(MESES_NOMES.get(mes_cod, mes_cod), style_td_bold),
            Paragraph(filial, style_td),
            Paragraph(format_number_br(qtd, 0), style_td_right),
            Paragraph(format_number_br(litros, 2), style_td_right),
            Paragraph(format_currency_br(valor), style_td_right_bold),
            Paragraph(f"R$ {p_medio:.3f}".replace('.', ','), style_td_right),
            Paragraph(format_currency_br(rec), style_td_right),
            Paragraph(format_currency_br(pend), style_td_right),
        ])

    t2 = Table(t2_data, colWidths=[105, 75, 80, 95, 110, 85, 110, 110])
    t2_style = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#2C3E50')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]
    for r_idx in range(1, len(t2_data)):
        if r_idx % 2 == 0:
            t2_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), colors.HexColor('#F8FAFC')))
    t2.setStyle(TableStyle(t2_style))
    story.append(t2)
    story.append(Spacer(1, 10))

    # Tabela 3: Faturamento por Segmento
    story.append(Paragraph("3. Faturamento por Segmento: Frota Própria (Consumo VIPE) vs Terceiros / Parceiros", style_section))
    t3_data = [
        [
            Paragraph("Mês", style_th),
            Paragraph("Segmento", style_th),
            Paragraph("Abastecimentos", style_th),
            Paragraph("Volume (Litros)", style_th),
            Paragraph("Faturamento (R$)", style_th),
            Paragraph("Preço Médio/L", style_th),
            Paragraph("Recebido (R$)", style_th),
            Paragraph("Pendente (R$)", style_th),
        ]
    ]
    for idx, r in enumerate(data['resumo_segmento']):
        mes_cod = r[0]
        segmento = r[1]
        qtd = r[2]
        litros = float(r[3])
        valor = float(r[4])
        p_medio = float(r[5])
        rec = float(r[6])
        pend = float(r[7])

        t3_data.append([
            Paragraph(MESES_NOMES.get(mes_cod, mes_cod), style_td_bold),
            Paragraph(segmento, style_td),
            Paragraph(format_number_br(qtd, 0), style_td_right),
            Paragraph(format_number_br(litros, 2), style_td_right),
            Paragraph(format_currency_br(valor), style_td_right_bold),
            Paragraph(f"R$ {p_medio:.3f}".replace('.', ','), style_td_right),
            Paragraph(format_currency_br(rec), style_td_right),
            Paragraph(format_currency_br(pend), style_td_right),
        ])

    t3 = Table(t3_data, colWidths=[95, 130, 85, 80, 105, 75, 100, 100])
    t3_style = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1E293B')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]
    for r_idx in range(1, len(t3_data)):
        if r_idx % 2 == 0:
            t3_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), colors.HexColor('#F8FAFC')))
    t3.setStyle(TableStyle(t3_style))
    story.append(t3)

    # PAGE BREAK: Página 3 para Top Proprietários
    story.append(PageBreak())

    story.append(Paragraph("4. Ranking dos Maiores Proprietários / Clientes (Top 15)", style_section))
    t4_data = [
        [
            Paragraph("Pos.", style_th),
            Paragraph("Proprietário / Cliente", style_th),
            Paragraph("Abastecimentos", style_th),
            Paragraph("Volume Total (L)", style_th),
            Paragraph("Faturamento Total (R$)", style_th),
            Paragraph("Part. %", style_th),
            Paragraph("Recebido (R$)", style_th),
            Paragraph("Pendente (R$)", style_th),
        ]
    ]

    for idx, r in enumerate(data['top_proprietarios'][:15], 1):
        p_nome = r[0]
        qtd = r[1]
        litros = float(r[2])
        valor = float(r[3])
        rec = float(r[4])
        pend = float(r[5])
        part_pct = (valor / tot_val * 100.0) if tot_val > 0 else 0

        t4_data.append([
            Paragraph(f"{idx}º", style_td_bold),
            Paragraph(p_nome, style_td_bold),
            Paragraph(format_number_br(qtd, 0), style_td_right),
            Paragraph(format_number_br(litros, 2), style_td_right),
            Paragraph(format_currency_br(valor), style_td_right_bold),
            Paragraph(f"{part_pct:.1f}%".replace('.', ','), style_td_right),
            Paragraph(format_currency_br(rec), style_td_right),
            Paragraph(format_currency_br(pend), style_td_right),
        ])

    t4 = Table(t4_data, colWidths=[35, 185, 80, 95, 115, 60, 100, 100])
    t4_style = [
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#1B365D')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 3.5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3.5),
    ]
    for r_idx in range(1, len(t4_data)):
        if r_idx % 2 == 0:
            t4_style.append(('BACKGROUND', (0, r_idx), (-1, r_idx), colors.HexColor('#F8FAFC')))
    t4.setStyle(TableStyle(t4_style))
    story.append(t4)
    story.append(Spacer(1, 14))

    # Observações finais
    obs_text = (
        "<b>Notas Explicativas:</b><br/>"
        "1. <b>Outubro/2026*:</b> Mês em andamento apurado até 20/10/2026.<br/>"
        "2. <b>Frota Própria (VIPE):</b> Refere-se às operações de abastecimento dos caminhões da VIPE TRANSPORTES (consumo próprio operacional).<br/>"
        "3. <b>Terceiros / Parceiros:</b> Representa o faturamento com transportadores terceirizados e agregados atendidos no ponto de abastecimento de Viana.<br/>"
        "4. A planilha detalhada com todos os 4.396 lançamentos individuais encontra-se disponível no arquivo <b>Relatorio_Faturamento_Mensal.xlsx</b>."
    )
    story.append(Paragraph(obs_text, style_subtitle))

    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"PDF salvo com sucesso em: {output_path}")


# ==============================================================================
# MAIN EXECUTION
# ==============================================================================
if __name__ == '__main__':
    root_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    excel_path = os.path.join(root_dir, "Relatorio_Faturamento_Mensal.xlsx")
    pdf_path = os.path.join(root_dir, "Relatorio_Faturamento_Mensal.pdf")
    charts_dir = os.path.join(root_dir, "tmp", "charts")

    print("Conectando ao banco de dados e obtendo dados operacionais...")
    data = fetch_all_data()
    print(f"Dados obtidos! Total abastecimentos: {len(data['base_detalhada'])}")

    print("Gerando planilha Excel (XLSX)...")
    gerar_excel(data, excel_path)

    print("Gerando relatório executivo PDF...")
    gerar_pdf(data, pdf_path, charts_dir)

    print("Todos os relatórios foram gerados com sucesso!")
